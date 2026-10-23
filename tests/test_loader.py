import pymupdf
import pytest
from docx import Document

from backend.ingestion.loader import extract, safe_filename


def test_pdf_pages(tmp_path):
    path = tmp_path / "test.pdf"
    document = pymupdf.open()
    for text in ["First page", "Second page"]:
        document.new_page().insert_text((72, 72), text)
    document.save(path)
    document.close()
    sections = extract(path, ".pdf")
    assert [s.page_number for s in sections] == [1, 2]
    assert "Second page" in sections[1].text


def test_docx_table(tmp_path):
    path = tmp_path / "test.docx"
    document = Document()
    document.add_paragraph("Before")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Inside table"
    document.add_paragraph("After")
    document.save(path)
    text = extract(path, ".docx")[0].text
    assert text.index("Before") < text.index("Inside table") < text.index("After")


@pytest.mark.parametrize("suffix", [".txt", ".md"])
def test_text_and_empty(tmp_path, suffix):
    path = tmp_path / ("test" + suffix)
    path.write_text("hello")
    assert extract(path, suffix)[0].text == "hello"
    path.write_text(" ")
    with pytest.raises(ValueError):
        extract(path, suffix)


def test_safe_filename():
    assert safe_filename("../../guide.pdf") == "guide.pdf"
    assert safe_filename("C:\\private\\guide.pdf") == "guide.pdf"
    with pytest.raises(ValueError):
        safe_filename("malware.exe")


def test_corrupt_pdf(tmp_path):
    path = tmp_path / "bad.pdf"
    path.write_bytes(b"not a pdf")
    with pytest.raises(ValueError):
        extract(path, ".pdf")
